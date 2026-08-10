"""
Create PowerPoint Presentation for Project FORESIGHT
Generates a professional 15-slide presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def create_presentation():
    """Create the PowerPoint presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define colors
    title_color = RgbColor(31, 119, 180)  # Blue
    text_color = RgbColor(50, 50, 50)     # Dark_gray
    accent_color = RgbColor(255, 127, 14) # Orange
    
    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    add_title_slide(slide, "Project FORESIGHT", 
                   "Demand Forecasting & Inventory Risk Prediction System",
                   "Intelligent Inventory Management Solution")
    
    # Slide 2: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Problem Statement", [
        "Businesses struggle with inventory optimization",
        "Stockouts lead to lost sales and customer dissatisfaction",
        "Overstocking ties up working capital and increases holding costs",
        "Manual inventory management is reactive, not proactive",
        "Lack of data-driven decision making in supply chain",
        "Need for accurate demand forecasting and risk prediction"
    ], title_color, text_color)
    
    # Slide 3: Project Objective
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Project Objective", [
        "Build an intelligent demand forecasting system",
        "Predict future demand using machine learning models",
        "Identify stockout and overstock risks in real-time",
        "Provide actionable recommendations for inventory decisions",
        "Enable data-driven inventory optimization",
        "Reduce costs and improve supply chain efficiency"
    ], title_color, text_color)
    
    # Slide 4: Solution Overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Solution Overview", [
        "End-to-end data processing pipeline",
        "Multiple ML forecasting models (Random Forest, LightGBM, Prophet)",
        "Intelligent risk prediction engine",
        "Multi-interface access: CLI, REST API, Interactive Dashboard",
        "Real-time inventory health monitoring",
        "Automated action recommendations"
    ], title_color, text_color)
    
    # Slide 5: System Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_architecture_slide(slide, "System Architecture", [
        "Data Layer: Raw CSV files (sales, inventory, calendar, SKU master)",
        "Processing Layer: Data cleaning, merging, feature engineering",
        "ML Layer: Model training, forecasting, risk prediction",
        "API Layer: REST endpoints for programmatic access",
        "Presentation Layer: Streamlit interactive dashboard",
        "Output Layer: CSV reports, visualizations, recommendations"
    ], title_color, text_color)
    
    # Slide 6: Data Pipeline
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Data Pipeline", [
        "Load raw data from CSV files",
        "Clean and validate data (remove duplicates, handle missing values)",
        "Merge datasets (sales, SKU master, calendar, inventory)",
        "Feature engineering:",
        "  - Temporal features (day, week, month, weekend)",
        "  - Lag features (previous sales)",
        "  - Rolling averages (7-day, 14-day)",
        "  - External factors (holidays, promotions, price)",
        "Save processed data for ML models"
    ], title_color, text_color)
    
    # Slide 7: ML Forecasting Models
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "ML Forecasting Models", [
        "Random Forest (Default):",
        "  - Robust, handles non-linear relationships",
        "  - Ensemble of decision trees",
        "  - Good for general-purpose forecasting",
        "",
        "LightGBM:",
        "  - Gradient boosting framework",
        "  - Faster training speed",
        "  - Handles large datasets efficiently",
        "",
        "Prophet:",
        "  - Time-series specific model by Facebook",
        "  - Handles seasonality and holidays",
        "  - Interpretable parameters"
    ], title_color, text_color)
    
    # Slide 8: Risk Prediction Engine
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Risk Prediction Engine", [
        "Stockout Risk Calculation:",
        "  - Ratio of forecast to current stock",
        "  - High: forecast >= stock",
        "  - Medium: forecast >= 80% of stock",
        "  - Low: forecast >= 50% of stock",
        "",
        "Overstock Risk Calculation:",
        "  - Ratio of stock to forecast",
        "  - High: stock >= 1.5x forecast",
        "  - Medium: stock >= 1.2x forecast",
        "  - Low: stock >= 1.0x forecast"
    ], title_color, text_color)
    
    # Slide 9: Action Recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Action Recommendations", [
        "Reorder Now:",
        "  - High/Medium stockout risk detected",
        "  - Immediate action required to prevent lost sales",
        "",
        "Markdown/Clear:",
        "  - High overstock risk detected",
        "  - Reduce inventory through promotions or clearance",
        "",
        "Watch:",
        "  - Low risk levels",
        "  - Monitor closely for changes",
        "",
        "Healthy:",
        "  - No significant risks",
        "  - Inventory at optimal levels"
    ], title_color, text_color)
    
    # Slide 10: Dashboard Features
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Interactive Dashboard", [
        "Overview Tab: Key metrics, risk distribution charts",
        "Forecasts Tab: Demand predictions, stock vs forecast comparison",
        "Risk Analysis Tab: Stockout/overstock breakdown, category health",
        "Recommendations Tab: Priority actions, color-coded table",
        "Product Details Tab: Individual SKU analysis, historical trends",
        "Real-time configuration: Model selection, threshold adjustment",
        "Interactive filtering: Search, category filters, risk filters"
    ], title_color, text_color)
    
    # Slide 11: REST API Integration
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "REST API Endpoints", [
        "POST /forecast - Generate demand forecasts",
        "GET /risk - Get risk summary statistics",
        "GET /recommendations - Get top priority actions",
        "GET /sku/{sku} - Get details for specific SKU",
        "GET /category/{category} - Get category-wise analysis",
        "GET /categories - List all available categories",
        "GET /skus - Get all SKUs with optional filters",
        "GET /health - API health check"
    ], title_color, text_color)
    
    # Slide 12: Technology Stack
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_tech_stack_slide(slide, "Technology Stack", [
        "Core: Python 3.13",
        "Data Processing: Pandas, NumPy",
        "ML Models: Scikit-learn, LightGBM, Prophet",
        "API Framework: FastAPI, Uvicorn",
        "Dashboard: Streamlit, Plotly",
        "Visualization: Matplotlib, Seaborn",
        "Development: JupyterLab, Notebook"
    ], title_color, text_color)
    
    # Slide 13: Key Benefits
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Key Business Benefits", [
        "Reduce inventory carrying costs",
        "Prevent lost sales from stockouts",
        "Optimize working capital allocation",
        "Improve supply chain efficiency",
        "Enable proactive inventory management",
        "Data-driven decision making",
        "Real-time risk monitoring",
        "Scalable and customizable solution"
    ], title_color, text_color)
    
    # Slide 14: Future Enhancements
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_content_slide(slide, "Future Enhancements", [
        "Additional ML models (XGBoost, ARIMA, LSTM)",
        "Multi-location inventory optimization",
        "Supplier lead time prediction",
        "Automated purchase order generation",
        "Integration with ERP systems",
        "Mobile app for on-the-go monitoring",
        "Advanced analytics and reporting",
        "Collaborative forecasting with team inputs"
    ], title_color, text_color)
    
    # Slide 15: Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_slide(slide, "Thank You", 
                   "Questions & Discussion",
                   "Project FORESIGHT - Intelligent Inventory Management")
    
    # Save presentation
    output_path = "Project_FORESIGHT_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")


def add_title_slide(slide, title, subtitle, footer):
    """Add title slide with centered text"""
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, slide.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(245, 245, 245)
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(
        Inches(1), Inches(2), Inches(8), Inches(1)
    )
    title_frame = title_box.text_frame
    title_frame.text = title
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(31, 119, 180)
    title_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide.shapes.add_textbox(
        Inches(1), Inches(3.5), Inches(8), Inches(1)
    )
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = subtitle
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = RgbColor(50, 50, 50)
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Footer
    footer_box = slide.shapes.add_textbox(
        Inches(1), Inches(5.5), Inches(8), Inches(0.5)
    )
    footer_frame = footer_box.text_frame
    footer_frame.text = footer
    footer_para = footer_frame.paragraphs[0]
    footer_para.font.size = Pt(18)
    footer_para.font.color.rgb = RgbColor(128, 128, 128)
    footer_para.alignment = PP_ALIGN.CENTER


def add_content_slide(slide, title, bullet_points, title_color, text_color):
    """Add content slide with title and bullet points"""
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, slide.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(255, 255, 255)
    background.line.fill.background()
    
    # Title bar
    title_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, Inches(1.2)
    )
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = RgbColor(31, 119, 180)
    title_bar.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(9), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = title
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(255, 255, 255)
    
    # Content box
    content_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.8), Inches(8.4), Inches(5)
    )
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    
    for i, point in enumerate(bullet_points):
        if i == 0:
            p = content_frame.paragraphs[0]
        else:
            p = content_frame.add_paragraph()
        
        p.text = point
        p.font.size = Pt(20)
        p.font.color.rgb = text_color
        p.space_after = Pt(10)
        
        if point and not point.startswith("  "):
            p.level = 0
        else:
            p.level = 1
            p.text = point.strip()


def add_architecture_slide(slide, title, components, title_color, text_color):
    """Add architecture diagram slide"""
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, slide.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(255, 255, 255)
    background.line.fill.background()
    
    # Title bar
    title_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, Inches(1.2)
    )
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = RgbColor(31, 119, 180)
    title_bar.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(9), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = title
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(255, 255, 255)
    
    # Architecture boxes
    colors = [
        RgbColor(255, 127, 14),   # Orange
        RgbColor(44, 160, 44),    # Green
        RgbColor(214, 39, 40),    # Red
        RgbColor(148, 103, 189),  # Purple
        RgbColor(140, 86, 75),    # Brown
        RgbColor(31, 119, 180)    # Blue
    ]
    
    y_positions = [1.5, 2.5, 3.5, 4.5, 5.5, 6.5]
    
    for i, (component, color) in enumerate(zip(components, colors)):
        # Box
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(y_positions[i]),
            Inches(8.4), Inches(0.6)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.color.rgb = RgbColor(50, 50, 50)
        box.line.width = Pt(1)
        
        # Text
        text_box = slide.shapes.add_textbox(
            Inches(1), Inches(y_positions[i] + 0.1),
            Inches(8), Inches(0.4)
        )
        text_frame = text_box.text_frame
        text_frame.text = component.split(":")[0] if ":" in component else component
        text_para = text_frame.paragraphs[0]
        text_para.font.size = Pt(16)
        text_para.font.bold = True
        text_para.font.color.rgb = RgbColor(255, 255, 255)
        text_para.alignment = PP_ALIGN.CENTER


def add_tech_stack_slide(slide, title, technologies, title_color, text_color):
    """Add technology stack slide with grid layout"""
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, slide.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(255, 255, 255)
    background.line.fill.background()
    
    # Title bar
    title_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, slide.slide_width, Inches(1.2)
    )
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = RgbColor(31, 119, 180)
    title_bar.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(9), Inches(0.6)
    )
    title_frame = title_box.text_frame
    title_frame.text = title
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(255, 255, 255)
    
    # Tech boxes in grid
    colors = [
        RgbColor(31, 119, 180),
        RgbColor(255, 127, 14),
        RgbColor(44, 160, 44),
        RgbColor(214, 39, 40),
        RgbColor(148, 103, 189),
        RgbColor(140, 86, 75),
        RgbColor(227, 119, 194)
    ]
    
    positions = [
        (0.8, 1.5), (4.8, 1.5),
        (0.8, 2.5), (4.8, 2.5),
        (0.8, 3.5), (4.8, 3.5),
        (2.8, 4.5)
    ]
    
    for i, (tech, color) in enumerate(zip(technologies, colors)):
        x, y = positions[i]
        
        # Box
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y),
            Inches(4), Inches(0.6)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.color.rgb = RgbColor(50, 50, 50)
        box.line.width = Pt(1)
        
        # Text
        text_box = slide.shapes.add_textbox(
            Inches(x + 0.2), Inches(y + 0.1),
            Inches(3.6), Inches(0.4)
        )
        text_frame = text_box.text_frame
        text_frame.text = tech
        text_para = text_frame.paragraphs[0]
        text_para.font.size = Pt(14)
        text_para.font.bold = True
        text_para.font.color.rgb = RgbColor(255, 255, 255)
        text_para.alignment = PP_ALIGN.CENTER


if __name__ == "__main__":
    create_presentation()
